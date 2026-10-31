# -*- coding: utf-8 -*-
"""Write a document.Document as an editable Word file (python-docx).

Right-to-left throughout: every Persian paragraph carries w:bidi and every
Persian run w:rtl with the complex-script font of the report settings (B Nazanin
by default); Latin runs are left-to-right in the Latin font. Tables are
w:bidiVisual with a repeating header row. Charts are 200 dpi PNG renderings of
the same figures as the HTML. The table of contents and page numbers are Word
fields; the file asks Word to update them when it is opened.

Alignment note: in a bidi paragraph Word treats 'left' and 'right' as logical,
so only start (no jc), 'both' and 'center' are used here."""

from __future__ import annotations

import io
import os
import re

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .. import APP_NAME, APP_VER
from ..core import fa_digits, fa_num, jalali_str
from ..render import to_png
from . import texts_fa as T
from .document import (Bullets, Callout, Cell, Cover, DocControl, FigureBlock,
                       FindingCard, Heading, PageBreak, Para, Signature, Table, TOC,
                       VerdictGrid, split_runs)

# characters XML 1.0 does not allow
_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

NAVY, NAVY2, GOLD, GOLD2 = "0F2742", "1D3D63", "C8902F", "E7C27D"
INK, INK2, MUTED, LINE, ZEBRA, HEAD = "16202B", "4B5563", "8A94A3", "E3E7ED", "F7F9FB", "EEF2F7"
TINT = {"ok": ("E5F4E5", "0B8A0B"), "warn": ("FFF3D1", "8A5E00"),
        "minor": ("FFF3D1", "8A5E00"), "material": ("FDE7DD", "B5532C"),
        "critical": ("FBE1E1", "B42424"), "fail": ("FBE1E1", "B42424"),
        "muted": ("EEF0F3", INK2), "info": ("E6F0FA", "1D5FA8")}
SOLID = {"ok": "0B8A0B", "warn": "B07A00", "material": "B5532C", "critical": "B42424",
         "muted": "6B7280"}
_VERDICT_TONE = {"SUITABLE": "ok", "RESTRICTED": "warn", "REMEDIATE": "critical"}
_SEV_TONE = {"CRITICAL": "critical", "MATERIAL": "material", "MINOR": "warn"}
CONTENT_W = 178.0                       # mm between the side margins


# ---------------------------------------------------------------------------
#  low-level XML helpers (children are written in schema order)
# ---------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _set_rpr(run, latin, bold=False, size=None, color=None, font_fa="B Nazanin",
             font_latin="Times New Roman", italic=False):
    r = run._r
    old = r.find(qn("w:rPr"))
    if old is not None:
        r.remove(old)
    rpr = _el("w:rPr")
    rpr.append(_el("w:rFonts", **{"w:ascii": font_latin, "w:hAnsi": font_latin,
                                  "w:cs": font_fa, "w:eastAsia": font_latin}))
    if bold:
        rpr.append(_el("w:b"))
        rpr.append(_el("w:bCs"))
    if italic:
        rpr.append(_el("w:i"))
        rpr.append(_el("w:iCs"))
    if color:
        rpr.append(_el("w:color", **{"w:val": color}))
    if size:
        half = str(int(round(size * 2)))
        rpr.append(_el("w:sz", **{"w:val": half}))
        rpr.append(_el("w:szCs", **{"w:val": half}))
    if not latin:
        rpr.append(_el("w:rtl"))
    rpr.append(_el("w:lang", **{"w:val": "en-US", "w:bidi": "fa-IR"}))
    r.insert(0, rpr)


_PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                   "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                   "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                   "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                   "w:pPrChange")


def _bidi(p, on=True):
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:bidi")):
        ppr.remove(old)
    if on:
        ppr.insert_element_before(_el("w:bidi"), *_PPR_AFTER_BIDI)


def _jc(p, val):
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:jc")):
        ppr.remove(old)
    if val:
        ppr.insert_element_before(_el("w:jc", **{"w:val": val}),
                                  "w:textDirection", "w:textAlignment",
                                  "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                                  "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _spacing(p, before=0, after=4, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    tcpr.insert_element_before(_el("w:shd", **{"w:val": "clear", "w:color": "auto",
                                               "w:fill": fill}),
                               "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                               "w:vAlign", "w:hideMark", "w:headers", "w:cellIns",
                               "w:cellDel", "w:cellMerge", "w:tcPrChange")


def _cell_margins(table, top=40, bottom=40, side=80):
    tblpr = table._tbl.tblPr
    mar = _el("w:tblCellMar")
    for side_name, v in (("w:top", top), ("w:start", side), ("w:left", side),
                         ("w:bottom", bottom), ("w:end", side), ("w:right", side)):
        mar.append(_el(side_name, **{"w:w": v, "w:type": "dxa"}))
    tblpr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription",
                                "w:tblPrChange")


def _borders(table, color=LINE, inside_h=True, box=False, size=4):
    tblpr = table._tbl.tblPr
    b = _el("w:tblBorders")
    for name in ("w:top", "w:left", "w:bottom", "w:right", "w:insideH", "w:insideV"):
        show = (name in ("w:top", "w:bottom", "w:left", "w:right") and box) or \
               (name == "w:insideH" and inside_h) or (name == "w:bottom")
        b.append(_el(name, **{"w:val": "single" if show else "nil", "w:sz": size,
                              "w:space": 0, "w:color": color}))
    tblpr.insert_element_before(b, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _rtl_table(table):
    tblpr = table._tbl.tblPr
    tblpr.insert_element_before(_el("w:bidiVisual"), "w:tblStyleRowBandSize",
                                "w:tblStyleColBandSize", "w:tblW", "w:jc",
                                "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _widths(table, mm_list):
    tw = int(sum(mm_list) * 56.6929)
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblpr.insert_element_before(_el("w:tblW", **{"w:w": tw, "w:type": "dxa"}),
                                "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
                                "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    tblpr.insert_element_before(_el("w:tblLayout", **{"w:type": "fixed"}),
                                "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                "w:tblDescription", "w:tblPrChange")
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), mm_list):
        gc.set(qn("w:w"), str(int(w * 56.6929)))
    for row in table.rows:
        for c, w in zip(row.cells, mm_list):
            c.width = Mm(w)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(_el("w:tblHeader"))
    trpr.append(_el("w:cantSplit"))


def _field(p, instr, placeholder, writer, size=None, color=None):
    r1 = p.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = p.add_run()
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = p.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = p.add_run(placeholder)
    writer.fmt(r4, False, size=size, color=color)
    r5 = p.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


# ===========================================================================
class _DW:
    def __init__(self, doc):
        self.doc = doc
        self.meta = doc.meta
        self.font_fa = self.meta.font_fa or "B Nazanin"
        self.font_latin = self.meta.font_latin or "Times New Roman"
        self.d = DocxDocument()
        self._page()
        self._styles()
        self._header_footer()
        s = self.d.settings.element
        s.insert_element_before(_el("w:updateFields", **{"w:val": "true"}),
                                "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                                "w:compat", "w:docVars", "w:rsids", "w:attachedSchema",
                                "w:themeFontLang", "w:clrSchemeMapping",
                                "w:doNotIncludeSubdocsInStats",
                                "w:doNotAutoCompressPictures", "w:forceUpgrade",
                                "w:captions", "w:readModeInkLockDown", "w:smartTagType",
                                "w:shapeDefaults", "w:doNotEmbedSmartTags",
                                "w:decimalSymbol", "w:listSeparator")
        zoom = s.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")     # required by the schema

    # -- setup -------------------------------------------------------------
    def _page(self):
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
        sec.left_margin = sec.right_margin = Mm(16)
        sec.header_distance = sec.footer_distance = Mm(9)
        sec._sectPr.insert_element_before(_el("w:bidi"), "w:rtlGutter", "w:docGrid",
                                          "w:printerSettings", "w:sectPrChange")

    def _style_font(self, style, size, color=None, bold=False):
        rpr = style.element.get_or_add_rPr()
        for tag in ("w:rFonts", "w:sz", "w:szCs", "w:color", "w:b", "w:bCs", "w:lang"):
            for old in rpr.findall(qn(tag)):
                rpr.remove(old)
        rpr.insert(0, _el("w:rFonts", **{"w:ascii": self.font_latin,
                                         "w:hAnsi": self.font_latin, "w:cs": self.font_fa,
                                         "w:eastAsia": self.font_latin}))
        style.font.size = Pt(size)
        rpr.append(_el("w:szCs", **{"w:val": str(int(size * 2))}))
        if color:
            style.font.color.rgb = None
            rpr.append(_el("w:color", **{"w:val": color}))
        if bold:
            style.font.bold = True
            rpr.append(_el("w:bCs"))
        rpr.append(_el("w:lang", **{"w:val": "en-US", "w:bidi": "fa-IR"}))

    def _styles(self):
        st = self.d.styles
        normal = st["Normal"]
        self._style_font(normal, 11)
        pf = normal.paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = 1.25
        ppr = normal.element.get_or_add_pPr()
        ppr.insert_element_before(_el("w:bidi"), *_PPR_AFTER_BIDI)
        for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY2),
                                  ("Heading 3", 11.5, INK)):
            s = st[name]
            self._style_font(s, size, color, bold=True)
            s.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 9)
            s.paragraph_format.space_after = Pt(5)
            s.paragraph_format.keep_with_next = True
            sppr = s.element.get_or_add_pPr()
            sppr.insert_element_before(_el("w:bidi"), *_PPR_AFTER_BIDI)

    def _header_footer(self):
        sec = self.d.sections[0]
        m = self.meta
        hp = sec.header.paragraphs[0]
        _bidi(hp)
        self.runs(hp, (f"{m.project} · " if m.project else "") + T.REPORT_TITLE
                  + (f"   |   {m.report_no}" if m.report_no else ""), size=8, color=MUTED)
        fp = sec.footer.paragraphs[0]
        _bidi(fp)
        _jc(fp, "center")
        self.runs(fp, "صفحهٔ ", size=8, color=MUTED)
        _field(fp, "PAGE", "۱", self, size=8, color=MUTED)

    # -- text ---------------------------------------------------------------
    def add_run(self, p, text):
        # control characters (a vertical tab pasted into a cell, say) are not
        # allowed in XML and would stop the whole Word file
        return p.add_run(_XML_BAD.sub("", text) if text else text)

    def fmt(self, run, latin, bold=False, size=None, color=None):
        _set_rpr(run, latin, bold, size, color, self.font_fa, self.font_latin)

    def runs(self, p, text, bold=False, size=None, color=None):
        for seg, latin, b in split_runs(text):
            self.fmt(self.add_run(p, seg), latin, bold or b, size, color)
        return p

    def para(self, container, text="", style=None, size=None, color=None, bold=False,
             align=None, after=4, before=0, ltr=False):
        """align None = start (right in RTL). Justify ('both') only for running
        prose: on a short line Arabic-script justification stretches words."""
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        _bidi(p, not ltr)
        _jc(p, "left" if ltr else align)
        _spacing(p, before, after)
        if text:
            if ltr:
                self.fmt(self.add_run(p, text), True, bold, size, color)
            else:
                self.runs(p, text, bold, size, color)
        return p

    def table(self, rows, cols, widths):
        t = self.d.add_table(rows=rows, cols=cols)
        _rtl_table(t)
        _widths(t, widths)
        return t

    def cell_text(self, cell, text, bold=False, size=9.5, color=INK, ltr=False,
                  align=None):
        p = cell.paragraphs[0]
        _bidi(p, not ltr)
        _jc(p, "left" if ltr else align)
        _spacing(p, 0, 0)
        p.paragraph_format.line_spacing = 1.1
        if ltr:
            self.fmt(self.add_run(p, text), True, bold, size - 0.5, color)
        else:
            self.runs(p, text, bold, size, color)
        return p

    # -- blocks -------------------------------------------------------------
    def cover(self, b):
        m = b.meta
        d = m.the_date()
        band = self.table(1, 1, [CONTENT_W])
        c = band.cell(0, 0)
        _shade(c, NAVY)
        self.cell_text(c, f"{APP_NAME} {APP_VER}", size=9, color=GOLD2, ltr=True)
        for _ in range(3):
            self.para(c, "", after=6)
        self.para(c, m.project or "", size=13, color=GOLD2, after=4)
        self.para(c, T.REPORT_TITLE, size=24, color="FFFFFF", bold=True, after=4)
        self.para(c, T.REPORT_SUBTITLE, size=12, color="C9D6E6", after=18)
        _borders(band, NAVY, inside_h=False, box=False, size=24)
        self.para(self.d, "", after=10)
        meta = [("کارفرما", m.client), ("مشاور", m.consultant), ("شمارهٔ گزارش", m.report_no),
                ("ویرایش", fa_digits(m.revision)),
                ("تاریخ", f"{jalali_str(d)} — {d.isoformat()}"), ("محل", m.location)]
        meta = [x for x in meta if x[1]]
        mt = self.table((len(meta) + 1) // 2, 2, [CONTENT_W / 2] * 2)
        _borders(mt)
        for i, (k, v) in enumerate(meta):
            cell = mt.cell(i // 2, i % 2)
            self.cell_text(cell, k, size=8.5, color=MUTED)
            self.para(cell, v, size=11.5, bold=True, after=2)
        self.para(self.d, "", after=8)
        order = ["REMEDIATE", "RESTRICTED", "SUITABLE"]
        worst = min((c_ for _, _, c_ in b.verdicts), key=order.index, default="SUITABLE")
        n_worst = sum(1 for _, _, c_ in b.verdicts if c_ == worst)
        status = (T.VERDICT_FA[worst] if worst == "SUITABLE"
                  else f"{T.VERDICT_FA[worst]} برای {fa_num(n_worst, 0)} عنصر")
        avg = T.RATING_FA.get(b.rating, b.rating) + (
            f" ({fa_num(b.score, 0)} از ۱۰۰)" if b.score is not None else "")
        st = self.table(1, 1 + len(b.stats), [70] + [(CONTENT_W - 70) / max(1, len(b.stats))]
                        * len(b.stats))
        badge = st.cell(0, 0)
        _shade(badge, SOLID[_VERDICT_TONE[worst]])
        self.cell_text(badge, "وضعیت کلی", size=8.5, color="FFFFFF")
        self.para(badge, status, size=14, bold=True, color="FFFFFF", after=2)
        self.para(badge, f"میانگین وزنی رتبه‌ها: {avg}", size=8.5, color="FFFFFF", after=2)
        for j, (k, v) in enumerate(b.stats, start=1):
            cell = st.cell(0, j)
            self.cell_text(cell, v, bold=True, size=16, color=NAVY, align="center")
            self.para(cell, k, size=8.5, color=MUTED, align="center", after=2)
        _borders(st, "FFFFFF", inside_h=False)
        self.para(self.d, "حکم کفایت داده به تفکیک عنصر", style="Heading 3", before=14)
        chips = self.table(1, len(b.verdicts) or 1,
                           [CONTENT_W / max(1, len(b.verdicts))] * max(1, len(b.verdicts)))
        for j, (el, name, code) in enumerate(b.verdicts):
            cell = chips.cell(0, j)
            bg, fg = TINT[_VERDICT_TONE[code]]
            _shade(cell, bg)
            self.cell_text(cell, f"{name}: {T.VERDICT_FA[code]}", bold=True, size=9.5,
                           color=fg, align="center")
        _borders(chips, "FFFFFF", inside_h=False, size=12)
        self.para(self.d, "", after=60)
        self.para(self.d, f"{m.author}" + (f" — {m.author_title}" if m.author else "")
                  + f"      تهیه‌شده با نرم‌افزار {APP_NAME} {APP_VER}", size=8.5,
                  color=MUTED)
        self.page_break()

    def doccontrol(self, b):
        m = b.meta
        self.para(self.d, T.H["control"], style="Heading 2", before=0)
        rows = [("عنوان", T.REPORT_TITLE), ("پروژه", m.project), ("شمارهٔ گزارش", m.report_no),
                ("ویرایش", fa_digits(m.revision)), ("تاریخ", jalali_str(m.the_date()))]
        t = self.table(len(rows), 2, [40, CONTENT_W - 40])
        _borders(t)
        for i, (k, v) in enumerate(rows):
            _shade(t.cell(i, 0), HEAD)
            self.cell_text(t.cell(i, 0), k, bold=True, color=NAVY2)
            self.cell_text(t.cell(i, 1), v or "—")
        self.para(self.d, "", after=6)
        people = [("تهیه‌کننده", m.author, m.author_title), ("بازبین", m.reviewer, ""),
                  ("تأییدکننده", m.approver, "")]
        t = self.table(4, 4, [30, 50, 58, 40])
        _borders(t)
        for j, h in enumerate(["نقش", "نام", "سمت", "امضا"]):
            _shade(t.cell(0, j), HEAD)
            self.cell_text(t.cell(0, j), h, bold=True, color=NAVY2)
        for i, (r, n, ti) in enumerate(people, start=1):
            self.cell_text(t.cell(i, 0), r)
            self.cell_text(t.cell(i, 1), n or "—")
            self.cell_text(t.cell(i, 2), ti or "")

    def toc(self):
        self.para(self.d, T.H["toc"], style="Heading 2", before=16)
        p = self.para(self.d, "", after=6)
        _field(p, 'TOC \\o "1-2" \\h \\z \\u',
               "فهرست مطالب هنگام باز کردن فایل به‌روز می‌شود (یا کلید F9 را بزنید).",
               self, size=10, color=MUTED)

    def heading(self, b):
        style = {1: "Heading 1", 2: "Heading 2"}.get(b.level, "Heading 3")
        p = self.d.add_paragraph(style=style)
        _bidi(p)
        a = b.anchor or ""
        if b.level == 1 and (a in ("summary", "actions") or a.startswith(("el-", "app-"))):
            p.paragraph_format.page_break_before = True
        if b.number:
            self.runs(p, f"{b.number}  ", color=GOLD)
        self.runs(p, b.text)
        if b.level == 1:
            ppr = p._p.get_or_add_pPr()
            bdr = _el("w:pBdr")
            bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": 10, "w:space": 4,
                                          "w:color": GOLD}))
            ppr.insert_element_before(bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens",
                                      "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                                      "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                                      "w:bidi", *_PPR_AFTER_BIDI)

    def paragraph(self, b):
        size = {"lead": 12, "note": 9, "small": 9}.get(b.style, 11)
        color = INK2 if b.style == "note" else None
        self.para(self.d, b.text, size=size, color=color, align="both")

    def bullets(self, b):
        style = "List Number" if b.ordered else "List Bullet"
        for item in b.items:
            p = self.d.add_paragraph(style=style)
            _bidi(p)
            _jc(p, "both")
            _spacing(p, 0, 2)
            self.runs(p, item)

    def table_block(self, b):
        n = len(b.header)
        widths = b.widths or [1] * n
        tot = float(sum(widths))
        mm = [CONTENT_W * w / tot for w in widths]
        t = self.table(1 + len(b.rows), n, mm)
        _borders(t)
        _cell_margins(t, 30, 30, 70)
        size = 8.5 if b.compact else 9.5
        hdr = t.rows[0]
        _repeat_header(hdr)
        for j, h in enumerate(b.header):
            _shade(hdr.cells[j], HEAD)
            self.cell_text(hdr.cells[j], h, bold=True, size=size, color=NAVY2)
        for i, row in enumerate(b.rows, start=1):
            for j, c in enumerate(row):
                cell = t.cell(i, j)
                if i % 2 == 0:
                    _shade(cell, ZEBRA)
                if isinstance(c, Cell):
                    if c.tone:
                        bg, fg = TINT.get(c.tone, TINT["muted"])
                        _shade(cell, bg)
                        self.cell_text(cell, c.text, bold=True, size=size, color=fg,
                                       ltr=c.ltr)
                    else:
                        self.cell_text(cell, c.text, size=size, ltr=c.ltr)
                else:
                    self.cell_text(cell, str(c), size=size)
        if b.caption:
            p = self.para(self.d, "", after=10, before=3)
            self.runs(p, f"جدول {b.number}. ", bold=True, size=9, color=NAVY2)
            self.runs(p, b.caption, size=9, color=INK2)

    def figure(self, b):
        png = to_png(b.fig, "light", dpi=200)
        p = self.d.add_paragraph()
        _bidi(p)
        _jc(p, "center")
        _spacing(p, 6, 2)
        p.paragraph_format.keep_with_next = True
        w, _ = b.fig.size or (7.2, 3.6)
        p.add_run().add_picture(io.BytesIO(png), width=Mm(min(CONTENT_W, w * 25.4)))
        cp = self.para(self.d, "", after=10)
        self.runs(cp, f"شکل {b.number}. ", bold=True, size=9, color=NAVY2)
        self.runs(cp, b.caption, size=9, color=INK2)

    def callout(self, b):
        tone = {"info": "info", "ok": "ok", "warn": "warn", "critical": "critical"}[b.kind]
        bg, fg = TINT[tone]
        t = self.table(1, 1, [CONTENT_W])
        c = t.cell(0, 0)
        _shade(c, bg)
        _borders(t, fg, inside_h=False, box=True, size=6)
        self.cell_text(c, b.title, bold=True, size=10.5, color=fg)
        self.para(c, b.text, size=10, after=2, align="both")
        self.para(self.d, "", after=4)

    def finding(self, f):
        t = self.table(2, 1, [CONTENT_W])
        _borders(t, LINE, inside_h=True, box=True, size=6)
        head, body = t.cell(0, 0), t.cell(1, 0)
        _shade(head, ZEBRA)
        self.cell_text(head, f.title_fa, bold=True, size=10.5)
        tone = _SEV_TONE[f.severity]
        p = self.para(head, "", after=1)
        self.runs(p, T.SEVERITY_FA[f.severity], bold=True, size=9, color=TINT[tone][1])
        self.runs(p, f"   ·   {T.STAGE_FA[f.stage]}   ·   مسئول: "
                     f"{T.OWNER_FA.get(f.owner, f.owner)}   ·   "
                     f"{T.TIMING_FA.get(f.timing, f.timing)}", size=9, color=INK2)
        p = body.paragraphs[0]
        _bidi(p)
        _jc(p, "both")
        _spacing(p, 2, 2)
        self.runs(p, "شواهد: ", bold=True, size=10, color=NAVY2)
        self.runs(p, f.evidence_fa, size=10)
        p = self.para(body, "", after=2)
        self.runs(p, "علت محتمل: ", bold=True, size=10, color=NAVY2)
        self.runs(p, f.cause_fa, size=10)
        self.para(body, "راهکار:", bold=True, size=10, color=NAVY2, after=1)
        for i, a in enumerate(f.actions_fa, start=1):
            p = self.para(body, "", after=1)
            self.runs(p, f"{fa_digits(str(i))}. ", bold=True, size=10, color=GOLD)
            self.runs(p, a, size=10)
        self.para(self.d, "", after=4)

    def verdicts(self, b):
        per = 4
        items = b.items
        for k in range(0, len(items), per):
            chunk = items[k:k + per]
            t = self.table(1, len(chunk), [CONTENT_W / len(chunk)] * len(chunk))
            _borders(t, "FFFFFF", inside_h=False, size=12)
            for j, (name, el, code, counts) in enumerate(chunk):
                c = t.cell(0, j)
                bg, fg = TINT[_VERDICT_TONE[code]]
                _shade(c, bg)
                self.cell_text(c, name, bold=True, size=13, color=NAVY)
                self.para(c, el, size=8, color=MUTED, after=1, ltr=True)
                self.para(c, T.VERDICT_FA[code], bold=True, size=9.5, color=fg, after=1)
                self.para(c, f"بحرانی {fa_num(counts['CRITICAL'], 0)} · مهم "
                             f"{fa_num(counts['MATERIAL'], 0)} · جزئی "
                             f"{fa_num(counts['MINOR'], 0)}", size=8.5, color=INK2, after=2)
            self.para(self.d, "", after=4)

    def signature(self, m):
        people = [("تهیه‌کننده", m.author, m.author_title), ("بازبین", m.reviewer, ""),
                  ("تأییدکننده", m.approver, "")]
        t = self.table(1, 3, [CONTENT_W / 3] * 3)
        _borders(t, LINE, inside_h=False, box=True, size=6)
        for j, (r, n, ti) in enumerate(people):
            c = t.cell(0, j)
            self.cell_text(c, r, bold=True, size=9.5, color=NAVY2)
            self.para(c, n or " ", size=10.5, after=1)
            self.para(c, ti or " ", size=8.5, color=MUTED, after=36)
            self.para(c, "امضا و تاریخ", size=8, color=MUTED, after=2)

    def page_break(self):
        p = self.d.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)


def write_docx(doc, path):
    w = _DW(doc)
    for b in doc.blocks:
        if isinstance(b, Cover):
            w.cover(b)
        elif isinstance(b, DocControl):
            w.doccontrol(b)
        elif isinstance(b, TOC):
            w.toc()
        elif isinstance(b, Heading):
            w.heading(b)
        elif isinstance(b, Para):
            w.paragraph(b)
        elif isinstance(b, Bullets):
            w.bullets(b)
        elif isinstance(b, Table):
            w.table_block(b)
        elif isinstance(b, FigureBlock):
            w.figure(b)
        elif isinstance(b, Callout):
            w.callout(b)
        elif isinstance(b, FindingCard):
            w.finding(b.finding)
        elif isinstance(b, VerdictGrid):
            w.verdicts(b)
        elif isinstance(b, Signature):
            w.signature(b.meta)
        elif isinstance(b, PageBreak):
            w.page_break()
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    w.d.save(path)
    return path
